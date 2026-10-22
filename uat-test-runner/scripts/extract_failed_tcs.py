#!/usr/bin/env python3
"""
extract_failed_tcs.py — Trích xuất danh sách TC ID cần retest từ kết quả Playwright.
Hỗ trợ 3 scope theo ISTQB standard:
  - verify : chỉ TC bị fail (quick smoke)
  - module : TC fail + TẤT CẢ TC cùng module (DEFAULT — catch regression)
  - full   : toàn bộ TC trong uat.json

Usage:
    # Module scope (default) — recommended sau khi fix bug
    python extract_failed_tcs.py test-results/results_flat.json \\
        --uat-json docs/uat/<feature>/uat.json \\
        --scope module \\
        --output test-results/retest-tc-ids.txt

    # Chỉ verify fix (scope hẹp)
    python extract_failed_tcs.py test-results/results.json \\
        --scope verify

    # Full regression
    python extract_failed_tcs.py test-results/results_flat.json \\
        --uat-json docs/uat/<feature>/uat.json \\
        --scope full

    # In grep pattern thay vì list (dùng trong shell)
    python extract_failed_tcs.py results_flat.json \\
        --uat-json docs/uat/<feature>/uat.json --scope module --grep
"""
import argparse
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

FAILED_STATUSES = {"failed", "timedOut", "interrupted"}
# Khớp pattern với flatten_pw_results.py — hỗ trợ cả TC-001 lẫn TC-CASE-001
TC_ID_PATTERN = re.compile(r'\b(TC-[A-Z0-9]*-?\d+)\b', re.IGNORECASE)

# Column indexes (1-based, khớp với write_results_to_excel.py)
COL_TC_ID  = 1
COL_MODULE = 4


def extract_tc_id(title: str) -> str | None:
    m = TC_ID_PATTERN.search(title)
    return m.group(1).upper() if m else None


def get_failed_tc_ids(results_path: Path) -> set[str]:
    """Lấy tập TC ID bị fail từ Playwright results.json."""
    with open(results_path, "r", encoding="utf-8") as f:
        data = json.load(f)

    tests = data.get("tests", [])
    if not tests and "suites" in data:
        tests = _flatten_suites(data["suites"])

    failed: set[str] = set()
    for test in tests:
        status = test.get("status", "")
        if "results" in test:
            for r in test["results"]:
                if r.get("status", "") in FAILED_STATUSES:
                    status = r["status"]
                    break
        if status not in FAILED_STATUSES:
            continue
        # Ưu tiên field tc_id (results_flat.json) — title chỉ là fallback
        tc_id = str(test.get("tc_id", "") or "").strip().upper()
        if not tc_id:
            tc_id = extract_tc_id(test.get("title", ""))
        if tc_id:
            failed.add(tc_id)
    return failed


def get_all_tcs_from_uat_json(uat_json_path: Path) -> dict[str, str]:
    """
    Đọc uat.json (handoff schema từ uat-test-writer), trả về {tc_id: module}.
    Module resolution: field "module" → "journey_ref" → "" (fallback).
    Đây là nguồn ƯU TIÊN — khớp doctrine "Excel không bao giờ là input".
    """
    if not uat_json_path.exists():
        print(f"❌ uat.json không tồn tại: {uat_json_path}", file=sys.stderr)
        sys.exit(1)
    with open(uat_json_path, "r", encoding="utf-8") as f:
        data = json.load(f)
    tc_module_map: dict[str, str] = {}
    for tc in data.get("test_cases", []):
        tc_id = str(tc.get("tc_id", "") or tc.get("id", "")).strip().upper()
        if not tc_id:
            continue
        module = str(tc.get("module", "") or tc.get("journey_ref", "")).strip()
        tc_module_map[tc_id] = module
    return tc_module_map


def get_all_tcs_from_excel(excel_path: Path) -> dict[str, str]:
    """
    Đọc Excel UAT, trả về dict {tc_id: module}.
    Yêu cầu openpyxl.
    """
    if not HAS_OPENPYXL:
        print("❌ openpyxl chưa cài. Chạy: pip install openpyxl", file=sys.stderr)
        sys.exit(1)
    if not excel_path.exists():
        print(f"❌ Excel không tồn tại: {excel_path}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(excel_path, data_only=True)
    if "Test Cases" not in wb.sheetnames:
        print("❌ Sheet 'Test Cases' không tồn tại trong Excel.", file=sys.stderr)
        sys.exit(1)

    ws = wb["Test Cases"]
    tc_module_map: dict[str, str] = {}
    for row in range(4, ws.max_row + 1):
        tc_id_val = ws.cell(row=row, column=COL_TC_ID).value
        if not tc_id_val:
            continue
        tc_id = str(tc_id_val).strip().upper()
        module = str(ws.cell(row=row, column=COL_MODULE).value or "").strip()
        tc_module_map[tc_id] = module
    return tc_module_map


def build_retest_scope(
    failed_ids: set[str],
    tc_module_map: dict[str, str],
    scope: str,
) -> tuple[set[str], dict]:
    """
    Xây dựng tập TC cần retest theo scope.
    Trả về (retest_ids, info_dict).
    """
    if scope == "verify":
        return failed_ids, {
            "scope": "verify",
            "failed": len(failed_ids),
            "same_module": 0,
            "total": len(failed_ids),
            "modules": [],
        }

    if scope == "full":
        all_ids = set(tc_module_map.keys())
        return all_ids, {
            "scope": "full",
            "failed": len(failed_ids),
            "same_module": len(all_ids) - len(failed_ids),
            "total": len(all_ids),
            "modules": ["(toàn bộ)"],
        }

    # scope == "module" (DEFAULT)
    # Lấy modules của các TC fail
    affected_modules: set[str] = set()
    for tc_id in failed_ids:
        mod = tc_module_map.get(tc_id, "")
        if mod:
            affected_modules.add(mod)

    # Collect tất cả TC trong các module đó
    module_tcs: set[str] = set()
    for tc_id, mod in tc_module_map.items():
        if mod in affected_modules:
            module_tcs.add(tc_id)

    retest_ids = failed_ids | module_tcs
    same_module_extra = retest_ids - failed_ids

    return retest_ids, {
        "scope": "module",
        "failed": len(failed_ids),
        "same_module": len(same_module_extra),
        "total": len(retest_ids),
        "modules": sorted(affected_modules),
    }


def _flatten_suites(suites: list) -> list:
    tests = []
    for suite in suites:
        if "tests" in suite:
            tests.extend(suite["tests"])
        if "suites" in suite:
            tests.extend(_flatten_suites(suite["suites"]))
    return tests


def main():
    parser = argparse.ArgumentParser(
        description="Trích xuất TC ID cần retest (ISTQB module-scope)"
    )
    parser.add_argument("results_file", help="Đường dẫn tới results.json (Playwright JSON reporter)")
    parser.add_argument(
        "--uat-json",
        "-j",
        help="uat.json từ uat-test-writer (nguồn ƯU TIÊN cho module map — scope=module/full)",
    )
    parser.add_argument(
        "--excel-file",
        "-e",
        help="[LEGACY fallback] File Excel UAT — chỉ dùng khi không có --uat-json",
    )
    parser.add_argument(
        "--scope",
        choices=["verify", "module", "full"],
        default="module",
        help="Scope retest: verify=chỉ fail | module=fail+cùng module (default) | full=tất cả",
    )
    parser.add_argument("--output", "-o", help="Ghi danh sách TC ID ra file (1 ID/dòng)")
    parser.add_argument("--grep", action="store_true", help="In grep pattern (TC-001|TC-002|...) thay vì list")
    parser.add_argument("--layer", choices=["api", "e2e", "all"], default="all")
    args = parser.parse_args()

    results_path = Path(args.results_file)
    if not results_path.exists():
        print(f"❌ File không tồn tại: {results_path}", file=sys.stderr)
        sys.exit(1)

    # Bước 1: Lấy TC fail
    failed_ids = get_failed_tc_ids(results_path)

    if not failed_ids:
        print("✅ Không có TC nào fail — không cần retest.", file=sys.stderr)
        if args.output:
            Path(args.output).unlink(missing_ok=True)
        sys.exit(0)

    # Bước 2: Build scope — module map từ uat.json (ưu tiên) hoặc Excel (legacy)
    tc_module_map: dict[str, str] = {}
    if args.scope in ("module", "full"):
        if args.uat_json:
            tc_module_map = get_all_tcs_from_uat_json(Path(args.uat_json))
        elif args.excel_file:
            tc_module_map = get_all_tcs_from_excel(Path(args.excel_file))
        else:
            # Tự tìm uat.json theo convention docs/uat/*/uat.json
            candidates = sorted(Path("docs/uat").glob("*/uat.json")) if Path("docs/uat").exists() else []
            if len(candidates) == 1:
                print(f"ℹ️  Tự phát hiện module map: {candidates[0]}", file=sys.stderr)
                tc_module_map = get_all_tcs_from_uat_json(candidates[0])
            else:
                print(
                    f"⚠️  Thiếu --uat-json (hoặc --excel-file legacy) cho scope={args.scope}. "
                    "Fallback về scope=verify.",
                    file=sys.stderr,
                )
                args.scope = "verify"

    retest_ids, info = build_retest_scope(failed_ids, tc_module_map, args.scope)

    # Sắp xếp: số cuối trong ID (TC-001 → 1, TC-CASE-017 → 17), fallback alphabet
    def _sort_key(x: str):
        m = re.search(r"(\d+)$", x)
        return (0, int(m.group(1)), x) if m else (1, 0, x)

    sorted_ids = sorted(retest_ids, key=_sort_key)

    # Bước 3: Output
    if args.grep:
        print("|".join(sorted_ids))
    else:
        print("\n".join(sorted_ids))

    # In summary ra stderr
    print(
        f"\n📋 Retest scope={info['scope']}: {info['total']} TC "
        f"({info['failed']} fail + {info['same_module']} cùng module)",
        file=sys.stderr,
    )
    if info["modules"] and info["scope"] == "module":
        print(f"   Modules bị ảnh hưởng: {', '.join(info['modules'])}", file=sys.stderr)

    if args.output:
        out_path = Path(args.output)
        out_path.parent.mkdir(parents=True, exist_ok=True)
        with open(out_path, "w", encoding="utf-8") as f:
            f.write("\n".join(sorted_ids) + "\n")
        print(f"   → Đã ghi: {out_path}", file=sys.stderr)


if __name__ == "__main__":
    main()

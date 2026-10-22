#!/usr/bin/env python3
"""
read_excel_uat.py — Đọc file Excel UAT 4 sheets → JSON cho Playwright runner.

Usage:
    # Đọc đầy đủ
    python read_excel_uat.py uat_feature.xlsx

    # Chỉ lấy TC tương ứng bug đã Fixed (cho retest)
    python read_excel_uat.py uat_feature.xlsx --filter-bugs-status=Fixed

    # Chỉ lấy TC theo priority
    python read_excel_uat.py uat_feature.xlsx --filter-priority=P0

    # Output ra file thay vì stdout
    python read_excel_uat.py uat_feature.xlsx -o test-results/parsed.json
"""

import argparse
import json
import sys
from pathlib import Path
from openpyxl import load_workbook


# Column mapping cho sheet Test Cases (header ở row 3, data từ row 4)
TC_COLS = {
    "tc_id": 1,                # A
    "type": 2,                 # B
    "priority": 3,             # C
    "module": 4,               # D
    "title": 5,                # E
    "requirement": 6,          # F
    "technique": 7,            # G
    "risk_level": 8,           # H
    "preconditions": 9,        # I
    "steps": 10,               # J
    "test_data": 11,           # K
    "data_required": 12,       # L — fixture IDs (CSV)
    "expected_ui": 13,         # M
    "expected_api": 14,        # N
    "expected_side": 15,       # O
    "severity_if_fail": 16,    # P
    "actual_result": 17,       # Q (runner-filled)
    "current_status_api": 18,  # R (runner-filled, layer 1)
    "current_status_e2e": 19,  # S (runner-filled, layer 2)
    "current_bug_id": 20,      # T (runner-filled)
    "last_run": 21,            # U (runner-filled)
    "run_history": 22,         # V (runner-filled)
    "regression_flag": 23,     # W (runner-filled)
    "notes": 24,               # X
}

# Data Matrix sheet columns
DM_COLS = {
    "fixture_id": 1,
    "entity": 2,
    "state": 3,
    "lifecycle": 4,
    "shape": 5,
    "description": 6,
    "setup_type": 7,
    "setup_recipe": 8,
    "cleanup_recipe": 9,
    "outputs": 10,
    "used_by_tcs": 11,
}

BUG_COLS = {
    "bug_id": 1,
    "tc_id": 2,
    "description": 3,
    "severity": 4,
    "status": 5,
    "discovered_date": 6,
    "last_fix_date": 7,
    "retest_result": 8,
    "notes": 9,
}


def parse_test_cases(ws):
    """Parse sheet Test Cases (header row 3, data từ row 4)."""
    test_cases = []
    for row in range(4, ws.max_row + 1):
        tc_id = ws.cell(row=row, column=TC_COLS["tc_id"]).value
        if not tc_id:
            continue

        tc = {}
        for key, col in TC_COLS.items():
            val = ws.cell(row=row, column=col).value
            tc[key] = str(val) if val is not None else ""

        tc["row_index"] = row  # Lưu để runner ghi đè đúng vị trí
        test_cases.append(tc)

    return test_cases


def parse_test_data(ws):
    """Parse sheet Test Data: 3 section accounts / boundary / payloads."""
    result = {"accounts": [], "boundary_values": [], "payloads": []}

    current_section = None
    current_headers = None

    for row in range(1, ws.max_row + 1):
        cell_a = ws.cell(row=row, column=1).value

        if not cell_a:
            continue

        cell_a_str = str(cell_a).strip()

        # Detect section
        if "TÀI KHOẢN" in cell_a_str.upper():
            current_section = "accounts"
            current_headers = None
            continue
        elif "BOUNDARY" in cell_a_str.upper():
            current_section = "boundary_values"
            current_headers = None
            continue
        elif "PAYLOAD" in cell_a_str.upper() or "SECURITY" in cell_a_str.upper():
            current_section = "payloads"
            current_headers = None
            continue

        if current_section is None:
            continue

        # Detect header row
        if current_headers is None:
            row_vals = [ws.cell(row=row, column=c).value for c in range(1, 8)]
            if any(v and isinstance(v, str) for v in row_vals):
                current_headers = row_vals
            continue

        # Parse data row
        item = {}
        for col_idx, header in enumerate(current_headers, start=1):
            if header:
                val = ws.cell(row=row, column=col_idx).value
                key = str(header).lower().strip().replace(" ", "_").replace("/", "_")
                item[key] = str(val) if val is not None else ""

        if any(item.values()):
            result[current_section].append(item)

    return result


def parse_bugs(ws):
    """Parse sheet Bug Tracker."""
    bugs = []
    for row in range(4, ws.max_row + 1):
        bug_id = ws.cell(row=row, column=BUG_COLS["bug_id"]).value
        if not bug_id:
            continue

        bug = {}
        for key, col in BUG_COLS.items():
            val = ws.cell(row=row, column=col).value
            bug[key] = str(val) if val is not None else ""

        bug["row_index"] = row
        bugs.append(bug)

    return bugs


def parse_data_matrix(ws):
    """Parse sheet Data Matrix → list fixture objects."""
    fixtures = []
    for row in range(4, ws.max_row + 1):
        fixture_id = ws.cell(row=row, column=DM_COLS["fixture_id"]).value
        if not fixture_id:
            continue

        fixture = {
            "fixture_id": str(fixture_id).strip(),
            "entity": str(ws.cell(row=row, column=DM_COLS["entity"]).value or "").strip(),
            "state": str(ws.cell(row=row, column=DM_COLS["state"]).value or "").strip(),
            "lifecycle": str(ws.cell(row=row, column=DM_COLS["lifecycle"]).value or "").strip(),
            "shape": str(ws.cell(row=row, column=DM_COLS["shape"]).value or "").strip(),
            "description": str(ws.cell(row=row, column=DM_COLS["description"]).value or "").strip(),
            "setup_type": str(ws.cell(row=row, column=DM_COLS["setup_type"]).value or "").strip(),
        }

        # Parse JSON fields
        for json_field in ("setup_recipe", "cleanup_recipe", "outputs"):
            raw = ws.cell(row=row, column=DM_COLS[json_field]).value
            if raw:
                try:
                    fixture[json_field] = json.loads(str(raw))
                except json.JSONDecodeError:
                    fixture[json_field] = str(raw)  # Keep as string nếu không parse được
            else:
                fixture[json_field] = {}

        used_by = ws.cell(row=row, column=DM_COLS["used_by_tcs"]).value
        fixture["used_by_tcs"] = [
            tc.strip() for tc in str(used_by or "").split(",") if tc.strip()
        ]

        fixtures.append(fixture)

    return fixtures


def read_uat_excel(file_path):
    """Đọc toàn bộ file UAT Excel."""
    wb = load_workbook(file_path, data_only=True)  # data_only để lấy giá trị formula

    if "Test Cases" not in wb.sheetnames:
        raise ValueError("File Excel phải có sheet 'Test Cases'")

    test_cases = parse_test_cases(wb["Test Cases"])
    test_data = parse_test_data(wb["Test Data"]) if "Test Data" in wb.sheetnames else {}
    bugs = parse_bugs(wb["Bug Tracker"]) if "Bug Tracker" in wb.sheetnames else []
    data_fixtures = parse_data_matrix(wb["Data Matrix"]) if "Data Matrix" in wb.sheetnames else []

    # Build fixture lookup for TC linking
    fixtures_by_id = {fx["fixture_id"]: fx for fx in data_fixtures}

    # Parse data_required từ TC string → list, và link tới fixture objects
    for tc in test_cases:
        raw_req = tc.get("data_required", "")
        if isinstance(raw_req, str):
            req_list = [r.strip() for r in raw_req.split(",") if r.strip()]
        else:
            req_list = raw_req
        tc["data_required"] = req_list
        # Inject resolved fixtures cho convenient access
        tc["resolved_fixtures"] = [
            fixtures_by_id[fid] for fid in req_list if fid in fixtures_by_id
        ]

    # Feature name từ Title cell A1
    feature_name = ""
    if "Test Cases" in wb.sheetnames:
        title = wb["Test Cases"].cell(row=1, column=1).value
        if title and ":" in str(title):
            feature_name = str(title).split(":", 1)[1].strip()

    return {
        "feature_name": feature_name,
        "file_path": str(Path(file_path).resolve()),
        "test_cases": test_cases,
        "test_data": test_data,
        "data_fixtures": data_fixtures,
        "fixtures_by_id": fixtures_by_id,
        "bugs": bugs,
        "summary": {
            "total_tc": len(test_cases),
            "total_fixtures": len(data_fixtures),
            "total_bugs": len([b for b in bugs if b.get("status") not in ("Closed", "Wont Fix", "")]),
        },
    }


def filter_by_bug_status(data, status):
    """Lọc TC tương ứng bug có Status cụ thể (cho retest workflow)."""
    matching_tc_ids = set()
    for bug in data["bugs"]:
        if bug.get("status", "").strip().lower() == status.lower():
            tc_id = bug.get("tc_id", "").strip()
            if tc_id:
                matching_tc_ids.add(tc_id)

    data["test_cases"] = [tc for tc in data["test_cases"] if tc["tc_id"] in matching_tc_ids]
    data["filter_applied"] = f"bugs_status={status}"
    data["matching_tc_ids"] = list(matching_tc_ids)
    return data


def filter_by_priority(data, priority):
    """Lọc TC theo priority."""
    data["test_cases"] = [tc for tc in data["test_cases"] if tc.get("priority", "").upper() == priority.upper()]
    data["filter_applied"] = f"priority={priority}"
    return data


def filter_by_type(data, tc_type):
    """Lọc TC theo loại."""
    data["test_cases"] = [tc for tc in data["test_cases"] if tc.get("type", "").upper() == tc_type.upper()]
    data["filter_applied"] = f"type={tc_type}"
    return data


def main():
    parser = argparse.ArgumentParser(description="Đọc file Excel UAT → JSON")
    parser.add_argument("excel_file", help="Đường dẫn tới file .xlsx")
    parser.add_argument("--filter-bugs-status", help="Lọc TC tương ứng bug có Status (Fixed/Open/...)")
    parser.add_argument("--filter-priority", help="Lọc theo priority (P0/P1/P2/P3)")
    parser.add_argument("--filter-type", help="Lọc theo loại TC (GREEN/RED/...)")
    parser.add_argument("-o", "--output", help="File output JSON (mặc định: stdout)")
    parser.add_argument("--pretty", action="store_true", help="Format JSON đẹp")

    args = parser.parse_args()

    if not Path(args.excel_file).exists():
        print(f"❌ File không tồn tại: {args.excel_file}", file=sys.stderr)
        sys.exit(1)

    try:
        data = read_uat_excel(args.excel_file)
    except Exception as e:
        print(f"❌ Lỗi đọc file: {e}", file=sys.stderr)
        sys.exit(2)

    # Apply filters
    if args.filter_bugs_status:
        data = filter_by_bug_status(data, args.filter_bugs_status)
    if args.filter_priority:
        data = filter_by_priority(data, args.filter_priority)
    if args.filter_type:
        data = filter_by_type(data, args.filter_type)

    # Output
    indent = 2 if args.pretty else None
    json_str = json.dumps(data, ensure_ascii=False, indent=indent)

    if args.output:
        Path(args.output).parent.mkdir(parents=True, exist_ok=True)
        Path(args.output).write_text(json_str, encoding="utf-8")
        print(f"✅ Đã ghi: {args.output}", file=sys.stderr)
        print(f"   - Feature: {data.get('feature_name', 'N/A')}", file=sys.stderr)
        print(f"   - Test cases: {len(data['test_cases'])}", file=sys.stderr)
        print(f"   - Bugs: {len(data['bugs'])}", file=sys.stderr)
    else:
        print(json_str)


if __name__ == "__main__":
    main()
